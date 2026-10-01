"""Step 4b: Standardize units in the extracted data.

The LLM writes units in different ways (MW, GW, GWe, MWp, TW; "USD billion" vs
"billion USD"). This converts them to one unit per measure so charts add up:
  - electric capacity -> GW
  - energy generated  -> TWh
  - money             -> USD billion
Thermal capacity (GWth) and anything unrecognized are left as they are.

Reads data/output/*.csv, writes the cleaned CSVs back, updates PostgreSQL,
and creates data/output/market_data.xlsx for Power BI.

Run:  python -m src.clean
"""
import pandas as pd

from src import config

# unit (lowercase) -> (standard unit, multiplier)
CONVERSIONS = {
    # capacity -> GW
    "gw": ("GW", 1), "gwe": ("GW", 1), "gwp": ("GW", 1), "gwac": ("GW", 1), "gwdc": ("GW", 1),
    "mw": ("GW", 0.001), "mwe": ("GW", 0.001), "mwp": ("GW", 0.001),
    "kw": ("GW", 0.000001), "tw": ("GW", 1000),
    # energy -> TWh
    "twh": ("TWh", 1), "twh/yr": ("TWh", 1), "twh/year": ("TWh", 1), "twh per year": ("TWh", 1),
    "gwh": ("TWh", 0.001), "gwh/yr": ("TWh", 0.001), "pwh": ("TWh", 1000),
    # money -> USD billion
    "usd billion": ("USD billion", 1), "billion usd": ("USD billion", 1),
    "$ billion": ("USD billion", 1), "billion $": ("USD billion", 1), "us$ billion": ("USD billion", 1),
    "usd million": ("USD billion", 0.001), "million usd": ("USD billion", 0.001),
    "usd trillion": ("USD billion", 1000), "trillion usd": ("USD billion", 1000),
    # percentages
    "%": ("%", 1), "percent": ("%", 1),
}


def standardize(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    if "original_unit" not in df.columns:  # keep the true original on re-runs
        df["original_unit"] = df["unit"]
    df["value"] = pd.to_numeric(df["value"], errors="coerce").astype(float)
    keys = df["unit"].astype(str).str.strip().str.lower()
    mapped = keys.map(CONVERSIONS)
    known = mapped.notna()
    df.loc[known, "unit"] = mapped[known].str[0]
    df.loc[known, "value"] = df.loc[known, "value"] * mapped[known].str[1]
    df["value"] = df["value"].round(4)
    return df


# Checked in order: first keyword match wins (e.g. "offshore wind" before "wind",
# "pumped hydro storage" counts as storage, not hydro).
TECH_GROUPS = [
    ("Offshore wind", ["offshore"]),
    ("Wind", ["wind"]),
    ("Solar", ["solar", "photovolt", "pv", "csp"]),
    ("Storage", ["batter", "storage", "pumped"]),
    ("Hydro", ["hydro"]),
    ("Geothermal", ["geotherm"]),
    ("Bioenergy", ["bio", "ethanol", "biofuel", "biodiesel", "biomass", "biogas", "wood"]),
    ("Marine", ["tidal", "wave", "marine", "ocean"]),
    ("Nuclear", ["nuclear"]),
    ("Fossil fuels", ["coal", "gas", "oil", "fossil"]),
    ("All renewables", ["renewable", "clean", "total", "all"]),
]


def group_technology(name) -> str:
    t = str(name).lower()
    for group, keywords in TECH_GROUPS:
        if any(k in t for k in keywords):
            return group
    return "Other"


def standardize_technology(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    if "original_technology" not in df.columns:  # keep the true original on re-runs
        df["original_technology"] = df["technology"]
    df["technology"] = df["original_technology"].map(group_technology)
    return df


def main():
    facts_path = config.OUTPUT_DIR / "market_facts.csv"
    events_path = config.OUTPUT_DIR / "company_events.csv"
    facts = pd.read_csv(facts_path)
    events = pd.read_csv(events_path)

    before = facts["unit"].nunique()
    facts = standardize(facts)
    print(f"Units: {before} different -> {facts['unit'].nunique()} after standardizing")
    print(facts["unit"].value_counts().head(10).to_string())

    before = facts["technology"].nunique()
    facts = standardize_technology(facts)
    print(f"\nTechnologies: {before} different -> {facts['technology'].nunique()} groups")
    print(facts["technology"].value_counts().to_string())

    facts.to_csv(facts_path, index=False)

    if config.DATABASE_URL:
        from sqlalchemy import create_engine

        engine = create_engine(config.DATABASE_URL)
        facts.to_sql("market_facts", engine, if_exists="replace", index=False)
        events.to_sql("company_events", engine, if_exists="replace", index=False)
        print("PostgreSQL tables updated.")

    xlsx = config.OUTPUT_DIR / "market_data.xlsx"
    with pd.ExcelWriter(xlsx) as w:
        facts.to_excel(w, sheet_name="market_facts", index=False)
        events.to_excel(w, sheet_name="company_events", index=False)
    format_as_tables(xlsx, {"market_facts": facts, "company_events": events})
    print(f"Power BI file ready: {xlsx}")


def format_as_tables(path, sheets: dict) -> None:
    """Power BI (web) only reads Excel data formatted as a table, so add one per sheet."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = load_workbook(path)
    for name, df in sheets.items():
        ws = wb[name]
        ref = f"A1:{get_column_letter(max(len(df.columns), 1))}{max(len(df), 1) + 1}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    wb.save(path)


if __name__ == "__main__":
    main()